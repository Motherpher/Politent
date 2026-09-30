#!/usr/bin/env python3
import csv
import re
import sys
from pathlib import Path
from openpyxl import load_workbook

def norm(value):
    if value is None:
        return ""
    s = str(value).strip().lower()
    s = s.replace("å","a").replace("ä","a").replace("ö","o")
    return re.sub(r"[^a-z0-9]+","",s)

HEADER_CANDIDATES = {
    "target_code_2026": ["Valdistriktets kod 2026","Valdistriktskod 2026","Kod 2026"],
    "target_name_2026": ["Valdistriktets namn 2026","Valdistriktsnamn 2026","Valdistrikt 2026"],
    "municipality": ["Kommun"],
    "county": ["Län","Lan"],
    "comparability": ["Jämförbarhet","Jamforbarhet"],
    "source_code_2022_1": ["Valdistriktskod 1 2022","Valdistriktskod 2022","Kod 1 2022"],
    "source_code_2022_2": ["Valdistriktskod 2 2022","Kod 2 2022"],
}

def find_col(headers, candidates):
    nh = [norm(h) for h in headers]
    nc = [norm(c) for c in candidates]
    for i,h in enumerate(nh):
        if h in nc:
            return i
    for i,h in enumerate(nh):
        if len(h) < 4:
            continue
        for c in nc:
            if h.startswith(c) or c.startswith(h):
                return i
    return None

def build_mapping(headers):
    return {k: find_col(headers,v) for k,v in HEADER_CANDIDATES.items()}

src = Path(sys.argv[1])
dst = Path(sys.argv[2])
wb = load_workbook(src, read_only=True, data_only=True)

required = {'target_code_2026','comparability','source_code_2022_1'}
candidates=[]

for ws in wb.worksheets:
    max_scan=min(ws.max_row, 40)
    for rownum,row in enumerate(ws.iter_rows(min_row=1,max_row=max_scan,values_only=True),start=1):
        headers=["" if v is None else str(v).strip() for v in row]
        mapping=build_mapping(headers)
        if not all(mapping[k] is not None for k in required):
            continue
        distinct={v for v in mapping.values() if v is not None}
        score=sum(v is not None for v in mapping.values())
        if len(distinct) < 4:
            continue
        candidates.append((score,ws.max_row,len(distinct),ws,rownum,headers,mapping))

if not candidates:
    raise SystemExit("Could not locate a tabular comparison header with required distinct columns")

# Prefer the most complete header; then the sheet with the most data rows.
candidates.sort(key=lambda x:(x[0],x[1],x[2]),reverse=True)
score,_,_,sheet,header_row,headers,mapping=candidates[0]
print(f"Using sheet={sheet.title!r}, header_row={header_row}, rows={sheet.max_row}, matched_columns={score}")
print("Header mapping:", {k: headers[v] if v is not None else None for k,v in mapping.items()})

dst.parent.mkdir(parents=True, exist_ok=True)
count=0
with dst.open("w",newline="",encoding="utf-8") as fh:
    writer=csv.DictWriter(fh,fieldnames=list(mapping))
    writer.writeheader()
    for row in sheet.iter_rows(min_row=header_row+1,values_only=True):
        record={}
        for key,idx in mapping.items():
            record[key]="" if idx is None or idx>=len(row) or row[idx] is None else str(row[idx]).strip()
        code=record["target_code_2026"]
        # Canonical polling-district codes are numeric-like; skip notes/footers.
        if not code or not re.fullmatch(r"\d{6,12}(?:\.0)?",code):
            continue
        if code.endswith(".0"):
            code=code[:-2]
        record["target_code_2026"]=code
        for k in ('source_code_2022_1','source_code_2022_2'):
            if record[k].endswith('.0'):
                record[k]=record[k][:-2]
        writer.writerow(record)
        count+=1

if count < 1000:
    raise SystemExit(f"Only {count} district rows parsed from sheet {sheet.title!r}; refusing incomplete crosswalk")

print(f"Parsed {count} polling-district comparison rows -> {dst}")
