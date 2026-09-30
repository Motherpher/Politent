#!/usr/bin/env python3
import hashlib
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

src = Path(sys.argv[1])
out = Path(sys.argv[2]) if len(sys.argv) > 2 else None

sha256 = hashlib.sha256(src.read_bytes()).hexdigest()
wb = load_workbook(src, read_only=True, data_only=True)
report = {
    "file": src.name,
    "sha256": sha256,
    "sheets": []
}

for ws in wb.worksheets:
    preview = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15), values_only=True):
        preview.append(["" if v is None else str(v) for v in row[:30]])
    report['sheets'].append({
        "title": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "preview": preview
    })

text = json.dumps(report, ensure_ascii=False, indent=2)
if out:
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text + '\n', encoding='utf-8')
else:
    print(text)
